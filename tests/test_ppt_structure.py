"""Unit tests for PPTX native structural preview."""

import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from src.parser import ppt_structure


def _make_pptx(path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    title = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    title.text_frame.text = "标题"
    body = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(2))
    body.text_frame.text = "正文第一行"
    prs.save(str(path))


class PptStructureTests(unittest.TestCase):
    def test_preview_returns_png_and_counts(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "sample.pptx"
            _make_pptx(path)
            result = ppt_structure.preview(path, 1)
        self.assertEqual(result["slide"], 1)
        self.assertEqual(result["slide_count"], 1)
        self.assertGreaterEqual(result["node_count"], 1)
        self.assertGreater(len(result["png"]), 100)  # a real PNG payload
        self.assertFalse(result["has_edges"])
        self.assertEqual(result["mermaid"], "")

    def test_preview_rejects_out_of_range_slide(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "sample.pptx"
            _make_pptx(path)
            with self.assertRaises(ValueError):
                ppt_structure.preview(path, 2)


if __name__ == "__main__":
    unittest.main()
