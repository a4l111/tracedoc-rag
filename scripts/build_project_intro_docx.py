from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "PROJECT_INTRO.md"
OUTPUT = ROOT / "docs" / "TraceDoc_RAG_项目介绍.docx"

BODY_FONT = "Microsoft YaHei"
LATIN_FONT = "Aptos"
MONO_FONT = "Cascadia Mono"
ACCENT = "1F4E78"
LIGHT_BLUE = "DCE6F1"
PALE_BLUE = "F3F7FB"
LIGHT_GRAY = "D9D9D9"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LIGHT_GRAY, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name=BODY_FONT, size=10.5, bold=None, color="000000") -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    r_pr = run._element.get_or_add_rPr()
    fonts = r_pr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, fonts)
    fonts.set(qn("w:ascii"), LATIN_FONT if name == BODY_FONT else name)
    fonts.set(qn("w:hAnsi"), LATIN_FONT if name == BODY_FONT else name)
    fonts.set(qn("w:eastAsia"), name)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    set_run_font(run, size=9, color="666666")


def add_inline_markdown(paragraph, text: str, size=10.5, color="000000") -> None:
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\[[^]]+\]\([^)]+\))")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            set_run_font(run, size=size, color=color)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name=MONO_FONT, size=max(8.5, size - 0.5), color="333333")
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        else:
            label, url = re.match(r"\[([^]]+)\]\(([^)]+)\)", token).groups()
            run = paragraph.add_run(f"{label}（{url}）")
            set_run_font(run, size=size, color=color)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=size, color=color)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.space_after = Pt(4)

    title = styles["Title"]
    title.font.name = BODY_FONT
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for style_name, size, before, after in (
        ("Heading 1", 17, 13, 6),
        ("Heading 2", 13.5, 10, 5),
        ("Heading 3", 11.5, 8, 3),
    ):
        style = styles[style_name]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_cover(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(85)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("TraceDoc RAG 项目介绍")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(16)
    subtitle.paragraph_format.space_after = Pt(30)
    run = subtitle.add_run("本地多模态文档可追溯问答系统")
    set_run_font(run, size=16, bold=True, color="333333")

    summary = doc.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    summary.paragraph_format.left_indent = Cm(2.2)
    summary.paragraph_format.right_indent = Cm(2.2)
    summary.paragraph_format.space_after = Pt(38)
    add_inline_markdown(
        summary,
        "面向 PDF、DOCX、PPTX、XLSX、TXT 和图片的解析、三路混合检索、Qwen3 精排与 DeepSeek 引用问答。",
        size=11.5,
        color="444444",
    )

    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.autofit = False
    widths = (Cm(4.0), Cm(8.2))
    values = (
        ("文档用途", "课程作业介绍与演示说明"),
        ("当前状态", "提交版候选，已加入实测图示"),
        ("评测日期", "2026-10-30"),
        ("评测范围", "18 题原始、复核及新快照三套结果"),
    )
    for row, values_row in zip(meta.rows, values):
        for index, (cell, value) in enumerate(zip(row.cells, values_row)):
            cell.width = widths[index]
            set_cell_shading(cell, PALE_BLUE if index else LIGHT_BLUE)
            set_cell_border(cell)
            set_cell_margins(cell, top=130, bottom=130)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(value)
            set_run_font(run, size=10, bold=index == 0)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(42)
    run = note.add_run("原始标签、人工复核标签和新快照指标分别展示，来源见第 8 节。")
    set_run_font(run, size=9.5, color="666666")
    doc.add_page_break()


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows):
        for column_index in range(columns):
            cell = table.cell(row_index, column_index)
            value = values[column_index] if column_index < len(values) else ""
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                set_cell_shading(cell, ACCENT)
            elif row_index % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if column_index > 0 else WD_ALIGN_PARAGRAPH.LEFT
            add_inline_markdown(
                paragraph,
                value,
                size=9.2,
                color="FFFFFF" if row_index == 0 else "000000",
            )
            for run in paragraph.runs:
                run.bold = row_index == 0
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        values = [part.strip() for part in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", value or "") for value in values):
            rows.append(values)
        index += 1
    return rows, index


def add_code_block(doc: Document, code_lines: list[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F2F2F2")
    set_cell_border(cell, color="C8C8C8")
    set_cell_margins(cell, top=120, start=150, bottom=120, end=150)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    for index, line in enumerate(code_lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        set_run_font(run, name=MONO_FONT, size=8.5, color="222222")
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_markdown_body(doc: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    index = 0
    in_code = False
    code_lines: list[str] = []
    skipped_title = False
    current_heading = ""

    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not stripped:
            index += 1
            continue
        if stripped.startswith("# ") and not skipped_title:
            skipped_title = True
            index += 1
            continue
        if stripped.startswith(">"):
            # Cover metadata already presents the opening status block.
            index += 1
            continue
        image_match = re.fullmatch(r"!\[([^]]+)\]\(([^)]+)\)", stripped)
        if image_match:
            caption, relative_path = image_match.groups()
            image_path = (SOURCE.parent / relative_path).resolve()
            if not image_path.is_relative_to(SOURCE.parent.resolve()):
                raise ValueError("Report image must be inside docs")
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run().add_picture(str(image_path), width=Cm(15.4))
            caption_paragraph = doc.add_paragraph()
            caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline_markdown(caption_paragraph, caption, size=8.5, color="555555")
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and lines[index + 1].strip().startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(doc, rows)
            continue

        heading_match = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading_match:
            level = min(len(heading_match.group(1)) - 1, 3)
            current_heading = heading_match.group(2)
            if current_heading.startswith("12.1 正式带引用问答截图"):
                doc.add_page_break()
            doc.add_heading(current_heading, level=level)
            index += 1
            continue

        bullet_match = re.match(r"^[-*]\s+(.*)$", stripped)
        number_match = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        checkbox_match = re.match(r"^-\s+\[([ xX])\]\s+(.*)$", stripped)
        if checkbox_match:
            paragraph = doc.add_paragraph(style="List Bullet")
            mark = "☒" if checkbox_match.group(1).lower() == "x" else "☐"
            add_inline_markdown(paragraph, f"{mark} {checkbox_match.group(2)}")
        elif bullet_match:
            paragraph = doc.add_paragraph(style="List Bullet")
            if current_heading == "13. 提交物":
                paragraph.paragraph_format.keep_with_next = True
            add_inline_markdown(paragraph, bullet_match.group(1))
        elif number_match:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.74)
            paragraph.paragraph_format.first_line_indent = Cm(-0.5)
            add_inline_markdown(paragraph, f"{number_match.group(1)}. {number_match.group(2)}")
        else:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.first_line_indent = Cm(0.74)
            add_inline_markdown(paragraph, stripped)
        index += 1


def build(output: Path = OUTPUT) -> Path:
    markdown = SOURCE.read_text(encoding="utf-8")
    architecture_start = markdown.index("```text", markdown.index("## 3. 系统架构"))
    architecture_end = markdown.index("```", architecture_start + 3) + 3
    markdown = (markdown[:architecture_start]
                + "```text\n文档 → Markdown / manifest / images\n"
                  "→ 固定窗口或结构化 Chunk → 共享快照\n"
                  "→ BM25 + Qwen3 Dense + bge-m3\n"
                  "→ RRF Top 50 → Qwen3 Reranker Top 5\n"
                  "→ DeepSeek 带引用回答 → CLI / Web\n```"
                + markdown[architecture_end:])
    checklist_start = markdown.index("### 8.4 提交前核对")
    next_section = markdown.index("## 9. 安全、隐私与可复现性", checklist_start)
    markdown = markdown[:checklist_start] + markdown[next_section:]
    doc = Document()
    configure_styles(doc)

    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.35)
    section.right_margin = Cm(2.35)
    section.different_first_page_header_footer = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("TraceDoc RAG 项目介绍")
    set_run_font(run, size=8.5, color="777777")
    add_page_number(section.footer.paragraphs[0])

    core = doc.core_properties
    core.title = "TraceDoc RAG 项目介绍"
    core.subject = "本地多模态文档可追溯问答系统课程项目介绍"
    core.author = "TraceDoc RAG 项目组"
    core.keywords = "RAG, 多模态文档, Qwen3, DeepSeek, 检索增强生成"

    add_cover(doc)
    add_markdown_body(doc, markdown)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output


if __name__ == "__main__":
    try:
        output = build(Path(sys.argv[1]) if len(sys.argv) > 1 else OUTPUT)
    except Exception as exc:
        print(f"生成失败：{exc}", file=sys.stderr)
        raise
    print(output)
