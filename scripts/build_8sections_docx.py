"""Build the 8-section course-report docx from PROJECT_INTRO_8sections.md.

Reuses the styling helpers from build_project_intro_docx.py but writes the
8-section cover (name / student id / college / major) and a condensed
evaluation section with no specific question numbers.
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

import build_project_intro_docx as base


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "PROJECT_INTRO_8sections.md"
OUTPUT = ROOT / "docs" / "TraceDoc_RAG_项目介绍_29题八节版_精简版.docx"


def add_cover(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(85)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("TraceDoc RAG 项目介绍")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(16)
    subtitle.paragraph_format.space_after = Pt(30)
    run = subtitle.add_run("本地多模态文档可追溯问答系统")
    base.set_run_font(run, size=16, bold=True, color="333333")

    summary = doc.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    summary.paragraph_format.left_indent = Cm(2.2)
    summary.paragraph_format.right_indent = Cm(2.2)
    summary.paragraph_format.space_after = Pt(38)
    base.add_inline_markdown(
        summary,
        "面向 PDF、DOCX、PPTX、XLSX、TXT 和图片的解析、三路混合检索、Qwen3 精排与 DeepSeek 引用问答。",
        size=11.5,
        color="444444",
    )

    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.autofit = False
    widths = (Cm(4.0), Cm(8.2))
    values = (
        ("姓名", "张舒昱"),
        ("学号", "265178"),
        ("学院", "东蒙国际学院（苏州）"),
        ("专业", "人工智能"),
    )
    for row, values_row in zip(meta.rows, values):
        for index, (cell, value) in enumerate(zip(row.cells, values_row)):
            cell.width = widths[index]
            base.set_cell_shading(cell, base.PALE_BLUE if index else base.LIGHT_BLUE)
            base.set_cell_border(cell)
            base.set_cell_margins(cell, top=130, bottom=130)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(value)
            base.set_run_font(run, size=10, bold=index == 0)

    doc.add_page_break()


def build(output: Path = OUTPUT) -> Path:
    base.SOURCE = SOURCE  # resolve relative figure paths against docs/
    markdown = SOURCE.read_text(encoding="utf-8")

    doc = Document()
    base.configure_styles(doc)

    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.35)
    section.right_margin = Cm(2.35)
    section.different_first_page_header_footer = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("TraceDoc RAG 项目介绍")
    base.set_run_font(run, size=8.5, color="777777")
    base.add_page_number(section.footer.paragraphs[0])

    core = doc.core_properties
    core.title = "TraceDoc RAG 项目介绍"
    core.subject = "本地多模态文档可追溯问答系统课程项目介绍"
    core.author = "TraceDoc RAG 项目组"
    core.keywords = "RAG, 多模态文档, Qwen3, DeepSeek, 检索增强生成"

    add_cover(doc)
    base.add_markdown_body(doc, markdown)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output


if __name__ == "__main__":
    try:
        out = build(Path(sys.argv[1]) if len(sys.argv) > 1 else OUTPUT)
    except Exception as exc:
        print(f"生成失败：{exc}", file=sys.stderr)
        raise
    print(out)
