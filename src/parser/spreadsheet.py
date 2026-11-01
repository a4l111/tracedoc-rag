"""Cell-level spreadsheet snapshots for browsing and bounded computation.

Reads cached cell values and formula text without executing formulas, macros,
links, or any embedded code. The resulting ``spreadsheet.json`` artifact is the
single input to the table browser and the deterministic table calculator.
"""

from __future__ import annotations

import csv as _csv
import hashlib
import io
import json
import math
import re
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

MAX_ROWS = 5000
MAX_COLS = 100
MAX_CELLS = 100_000
MAX_TEXT = 10_000


def cell_value(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("表格包含非有限数值")
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def display_value(value, fmt="General"):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if re.fullmatch(r"0(?:\.[0#]+)?%", fmt):
            places = len(fmt.split(".")[1].rstrip("%")) if "." in fmt else 0
            return f"{value * 100:.{places}f}%"
        if re.fullmatch(r"0{2,}", fmt) and isinstance(value, int):
            return str(value).zfill(len(fmt))
        if re.fullmatch(r"[#0,]+(?:\.[#0]+)?", fmt):
            places = len(fmt.split(".")[1]) if "." in fmt else 0
            return format(value, ("," if "," in fmt else "") + "." + str(places) + "f")
    return str(cell_value(value))


def _check_text(value):
    if isinstance(value, str) and len(value) > MAX_TEXT:
        raise ValueError("单元格文字超过10000字符，请拆分后导入；未截断原数据")


def read_xlsx(source: Path) -> tuple[list[dict], list[str]]:
    """Return (sheets, warnings) with cell-level data; no formula evaluation."""
    warnings: list[str] = []
    sheets: list[dict] = []
    total_cells = 0
    formulas = load_workbook(source, read_only=True, data_only=False, keep_links=False)
    cached = load_workbook(source, read_only=True, data_only=True, keep_links=False)
    try:
        for name in formulas.sheetnames:
            formula_sheet = formulas[name]
            cached_sheet = cached[name]
            if formula_sheet.sheet_state != "visible":
                warnings.append(f"已跳过隐藏工作表“{name}”。")
                continue
            rows: list[dict] = []
            max_col = 0
            last_row = 0
            for row_number, (f_row, c_row) in enumerate(
                    zip(formula_sheet.iter_rows(), cached_sheet.iter_rows()), start=1):
                if row_number > MAX_ROWS:
                    warnings.append(f"工作表“{name}”仅读取前 {MAX_ROWS} 行。")
                    break
                last_row = row_number
                cells: list[dict] = []
                for column, (f_cell, c_cell) in enumerate(zip(f_row, c_row), start=1):
                    if column > MAX_COLS:
                        break
                    raw = f_cell.value
                    if raw is None and c_cell.value is None:
                        continue
                    formula = ""
                    if f_cell.data_type == "f":
                        formula = str(raw) if isinstance(raw, str) else "（复杂公式，未展开）"
                    saved = c_cell.value
                    value = cell_value(saved if formula else raw)
                    _check_text(value)
                    if formula:
                        _check_text(formula)
                    fmt = f_cell.number_format or "General"
                    shown = display_value(saved if formula else raw, fmt)
                    if formula:
                        shown = f"公式缓存（未重算）：{shown}" if saved is not None else "公式未计算"
                    cells.append({
                        "column": column,
                        "address": f"{get_column_letter(column)}{row_number}",
                        "value": value,
                        "display": shown,
                        "format": fmt,
                        "formula": formula,
                        "type": "formula" if formula else f_cell.data_type,
                    })
                    max_col = max(max_col, column)
                if cells:
                    rows.append({"number": row_number, "cells": cells})
                total_cells += len(cells)
                if total_cells > MAX_CELLS:
                    raise ValueError("工作簿非空单元格超过100000个，请拆分文件")
            sheets.append({"name": name, "max_row": last_row, "max_col": max_col, "rows": rows})
    except (ValueError, KeyError) as exc:
        raise ValueError("无法读取 Excel：文件可能损坏、已加密或不是有效 .xlsx，请解密并另存后导入") from exc
    finally:
        formulas.close()
        cached.close()
    if not sheets:
        raise ValueError("工作簿没有可见工作表，请先在表格软件中取消隐藏")
    return sheets, warnings


def read_csv(source: Path) -> tuple[list[dict], list[str]]:
    raw = source.read_bytes()
    encoding = "utf-8-sig"
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        encoding = "utf-16"
    try:
        text = raw.decode(encoding)
    except UnicodeDecodeError:
        try:
            text = raw.decode("gb18030")
            encoding = "gb18030"
        except UnicodeDecodeError:
            raise ValueError("CSV 编码无法识别，请另存为 UTF-8 CSV") from None
    if "\x00" in text:
        raise ValueError("CSV 包含无效控制字符")
    try:
        dialect = _csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except _csv.Error:
        dialect = _csv.excel
    rows: list[dict] = []
    max_col = 0
    count = 0
    total_text = 0
    _csv.field_size_limit(1024 * 1024)
    number = 0
    try:
        for number, values in enumerate(_csv.reader(io.StringIO(text, newline=""), dialect), start=1):
            if number > MAX_ROWS or len(values) > MAX_COLS:
                raise ValueError(f"当前支持 CSV 最多 {MAX_ROWS} 行、{MAX_COLS} 列，请拆分后导入")
            cells: list[dict] = []
            max_col = max(max_col, len(values))
            for col, value in enumerate(values, start=1):
                _check_text(value)
                total_text += len(value)
                if total_text > 5_000_000:
                    raise ValueError("CSV 文字总量过大，请拆分后导入")
                if value != "":
                    cells.append({"column": col, "address": f"{get_column_letter(col)}{number}",
                                  "value": value, "display": value, "format": "文本",
                                  "formula": "", "type": "s"})
            count += len(cells)
            if count > MAX_CELLS:
                raise ValueError("CSV 非空单元格超过100000个，请拆分文件")
            if cells:
                rows.append({"number": number, "cells": cells})
    except _csv.Error:
        raise ValueError("CSV 结构或单元格大小不合法，请检查引号及分隔符") from None
    sheet = {"name": "CSV", "max_row": number, "max_col": max_col, "rows": rows}
    return [sheet], [f"CSV 按 {encoding} 读取，单元格保留文本，不执行以等号开头的内容；行号为 CSV 记录序号。"]


def detect_header(sheet: dict) -> None:
    """Mark the first (of 20) fully-text row with at least two cells as the header."""
    sheet["header_row"] = None
    for row in sheet["rows"][:20]:
        if len(row["cells"]) < 2:
            continue
        if all(cell["type"] in {"s", "str", "inlineStr"} and not cell["formula"] for cell in row["cells"]):
            sheet["header_row"] = row["number"]
            return


def write_artifact(out_dir: Path, source: Path, source_hash: str, sheets: list[dict], warnings: list[str]) -> None:
    for sheet in sheets:
        detect_header(sheet)
    artifact = {"version": 1, "source_name": source.name, "source_hash": source_hash,
                "sheets": sheets, "warnings": warnings}
    (out_dir / "spreadsheet.json").write_text(json.dumps(artifact, ensure_ascii=False), encoding="utf-8")


def snapshot(parsed_dir: Path, doc_id: str) -> dict:
    """Load the cell-level snapshot plus its source metadata for browsing/compute."""
    artifact_path = parsed_dir / doc_id / "spreadsheet.json"
    manifest_path = parsed_dir / doc_id / "manifest.json"
    if not artifact_path.is_file():
        raise FileNotFoundError("该文档没有可用的表格快照，请重新解析 Excel/CSV 文件")
    artifact = json.loads(artifact_path.read_text(encoding="utf-8"))
    manifest = json.loads(manifest_path.read_text(encoding="utf-8")) if manifest_path.is_file() else {}
    return {"id": doc_id, "name": manifest.get("source_file", artifact.get("source_name", doc_id)),
            "hash": manifest.get("sha256") or artifact.get("source_hash"), "sheets": artifact["sheets"],
            "warnings": artifact.get("warnings", [])}
