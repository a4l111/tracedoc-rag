"""PPTX native shape/group geometry and explicit connector-endpoint preview.

The generated image is a labelled structural preview built from the slide XML
(shape positions, text and explicit connector endpoints), not an Office slide
renderer. Arrow direction is never inferred; connectors are marked undirected.
"""

from __future__ import annotations

import io
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

MAX_NODES = 200
MAX_EDGES = 500


def _walk(items, scale=(1.0, 1.0), offset=(0.0, 0.0), parent=""):
    """Flatten a slide's shape tree, resolving group transforms into absolute EMU bboxes."""
    output = []
    for shape in items:
        key = f"{parent}/{shape.shape_id}" if parent else str(shape.shape_id)
        bbox = [offset[0] + shape.left * scale[0], offset[1] + shape.top * scale[1],
                shape.width * scale[0], shape.height * scale[1]]
        if shape.shape_type == 6:  # GROUP
            xfrm = shape.element.grpSpPr.xfrm
            sx = bbox[2] / xfrm.chExt.cx if xfrm is not None and xfrm.chExt is not None and xfrm.chExt.cx else scale[0]
            sy = bbox[3] / xfrm.chExt.cy if xfrm is not None and xfrm.chExt is not None and xfrm.chExt.cy else scale[1]
            cx = xfrm.chOff.x if xfrm is not None and xfrm.chOff is not None else 0
            cy = xfrm.chOff.y if xfrm is not None and xfrm.chOff is not None else 0
            output.extend(_walk(shape.shapes, (sx, sy), (bbox[0] - cx * sx, bbox[1] - cy * sy), key))
        else:
            output.append({"shape": shape, "path": key, "bbox": bbox, "shape_id": str(shape.shape_id)})
    return output


def _quote(text: str) -> str:
    return "".join(f"#{ord(c)};" if c in '"<>[]{}|`\\&\n\r' else c for c in text)


def _mermaid(nodes: list[dict], edges: list[dict]) -> str:
    ids = {node["id"]: f"n{i}" for i, node in enumerate(nodes)}
    lines = ["flowchart TD"]
    for node in nodes:
        lines.append(f'  {ids[node["id"]]}["{_quote(node["label"])}"]')
    for edge in edges:
        label = _quote(edge.get("label", ""))
        arrow = f'---|"{label}"|' if label else "---"
        lines.append(f'  {ids[edge["source"]]} {arrow} {ids[edge["target"]]}')
    return "\n".join(lines)


def preview(source: Path, slide_number: int) -> dict:
    """Return a structural preview of one slide: PNG bytes, Mermaid, and a text summary."""
    presentation = Presentation(str(source))
    if slide_number < 1 or slide_number > len(presentation.slides):
        raise ValueError(f"幻灯片号 {slide_number} 超出范围 1-{len(presentation.slides)}")
    slide = presentation.slides[slide_number - 1]
    width = presentation.slide_width or 9144000
    height = presentation.slide_height or 6858000

    items = _walk(slide.shapes)
    nodes, edges, by_id = [], [], {}
    for item in items:
        shape = item["shape"]
        if shape.has_text_frame and shape.text.strip():
            node = {"id": item["path"], "label": shape.text.strip()[:2000]}
            nodes.append(node)
            by_id[item["shape_id"]] = node["id"]
    for item in items:
        shape = item["shape"]
        if shape.shape_type != 9:  # connector
            continue
        starts = shape.element.xpath(".//a:stCxn")
        ends = shape.element.xpath(".//a:endCxn")
        if starts and ends and starts[0].get("id") in by_id and ends[0].get("id") in by_id:
            edges.append({"source": by_id[starts[0].get("id")], "target": by_id[ends[0].get("id")],
                          "label": "原生连接；方向未确认"})

    has_edges = bool(edges) and len(nodes) <= MAX_NODES and len(edges) <= MAX_EDGES
    if has_edges:
        mermaid = _mermaid(nodes, edges)
    else:
        mermaid = ""

    # Render a labelled structural thumbnail.
    factor = min(1600 / max(width, 1), 1000 / max(height, 1))
    canvas = Image.new("RGB", (max(1, int(width * factor)), max(1, int(height * factor))), "white")
    draw = ImageDraw.Draw(canvas)
    font_path = Path("C:/Windows/Fonts/msyh.ttc")
    font = ImageFont.truetype(str(font_path), 18) if font_path.exists() else ImageFont.load_default()
    for item in items:
        shape = item["shape"]
        x, y, w, h = [int(v * factor) for v in item["bbox"]]
        if shape.shape_type == 9:
            draw.line((x, y, x + w, y + h), fill="#738094", width=2)
        elif shape.shape_type == 13:
            try:
                with Image.open(io.BytesIO(shape.image.blob)) as image:
                    image = image.convert("RGB")
                    image.thumbnail((max(1, w), max(1, h)))
                    canvas.paste(image, (x, y))
            except OSError:
                pass
        elif shape.has_text_frame and shape.text.strip():
            draw.rectangle((x, y, x + max(1, w), y + max(1, h)), outline="#6872b4", width=2)
            label = shape.text.strip()
            line, lines = "", []
            for char in label:
                if char == "\n" or draw.textlength(line + char, font=font) > max(20, w - 12):
                    lines.append(line)
                    line = ""
                if char != "\n":
                    line += char
            lines.append(line)
            draw.multiline_text((x + 5, y + 4), "\n".join(lines[:max(1, h // 23)]),
                                fill="#253354", font=font, spacing=3)
    raw = io.BytesIO()
    canvas.save(raw, format="PNG")

    description = ("由 PPTX 原生形状、组合层级和显式连接关系提取。此图为结构预览，不是完整幻灯片渲染；"
                   "箭头方向与复杂样式仍需核对原文件。")
    if has_edges:
        description += "\n\n```mermaid\n" + mermaid + "\n```\n\n"
        description += "[原生结构提取：连接端点来自文档 XML，未推断箭头方向；请核对原幻灯片。]"

    return {"slide": slide_number, "slide_count": len(presentation.slides),
            "png": raw.getvalue(), "mermaid": mermaid, "description": description,
            "node_count": len(nodes), "edge_count": len(edges), "has_edges": has_edges}
