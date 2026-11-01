"""Local document parsing and per-file output management."""

import csv as _csv
import hashlib
import json
import posixpath
import re
import traceback
import zipfile
from functools import lru_cache
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

import fitz
from charset_normalizer import from_bytes
from PIL import Image, UnidentifiedImageError
from openpyxl import load_workbook
from rapidocr_onnxruntime import RapidOCR

from src.parser.tables import docx_table, markdown_table
from src.parser.pptx import parse_pptx


EXTENSIONS = {".pdf", ".docx", ".pptx", ".xlsx", ".csv", ".txt", ".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tif", ".tiff"}
IMAGE_EXTENSIONS = EXTENSIONS - {".pdf", ".docx", ".pptx", ".xlsx", ".csv", ".txt"}


def supported_files(raw_dir: Path) -> list[Path]:
    return sorted(p for p in raw_dir.rglob("*") if p.is_file() and p.suffix.lower() in EXTENSIONS)


def _sha256(source: Path) -> str:
    digest = hashlib.sha256()
    with source.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _doc_id(source: Path, raw_dir: Path) -> str:
    relative = source.relative_to(raw_dir).as_posix()
    label = "".join(c if c.isalnum() or c in "-_" else "-" for c in source.stem).strip("-") or "document"
    return f"{label[:48]}-{hashlib.sha256(relative.encode('utf-8')).hexdigest()[:12]}"


def _write_json(path: Path, data: dict[str, Any]) -> None:
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def parse_file(source: Path, raw_dir: Path, output_dir: Path, render_pages: bool = False,
               describe_images: bool = False, ocr_scanned_pages: bool = False) -> bool:
    doc_id = _doc_id(source, raw_dir)
    target = output_dir / doc_id
    images_dir = target / "images"
    images_dir.mkdir(parents=True, exist_ok=True)
    manifest: dict[str, Any] = {
        "doc_id": doc_id,
        "source_file": source.relative_to(raw_dir).as_posix(),
        "source_type": source.suffix.lower().lstrip("."),
        "sha256": None,
        "page_map": [],
        "image_assets": [],
        "errors": [],
    }
    errors_path = target / "errors.jsonl"
    errors_path.write_text("", encoding="utf-8")
    try:
        manifest["sha256"] = _sha256(source)
        suffix = source.suffix.lower()
        if suffix == ".pdf":
            markdown = _parse_pdf(source, images_dir, manifest, render_pages, ocr_scanned_pages)
        elif suffix == ".docx":
            markdown = _parse_docx(source, images_dir, manifest)
        elif suffix == ".pptx":
            markdown = parse_pptx(source, images_dir, manifest, _ocr_text)
        elif suffix == ".xlsx":
            markdown = _parse_xlsx(source, manifest)
        elif suffix == ".csv":
            markdown = _parse_csv(source, manifest)
        elif suffix == ".txt":
            markdown = _parse_txt(source, manifest)
        elif suffix in IMAGE_EXTENSIONS:
            markdown = _parse_image(source, images_dir, manifest)
        else:
            raise ValueError(f"Unsupported file type: {suffix}")
        if suffix in {".xlsx", ".csv"}:
            _write_spreadsheet_snapshot(source, target, manifest)
        if describe_images and manifest["image_assets"]:
            from src.parser.vision import enrich_assets
            markdown = enrich_assets(markdown, target, manifest)
        (target / "document.md").write_text(markdown, encoding="utf-8")
        return True
    except Exception as exc:
        error = {"file": str(source), "type": type(exc).__name__, "message": str(exc)}
        manifest["errors"].append(error)
        error["traceback"] = traceback.format_exc()
        (target / "document.md").write_text("", encoding="utf-8")
        return False
    finally:
        with errors_path.open("w", encoding="utf-8") as stream:
            for error in manifest["errors"]:
                stream.write(json.dumps(error, ensure_ascii=False) + "\n")
        _write_json(target / "manifest.json", manifest)


def _parse_pdf(source: Path, images_dir: Path, manifest: dict[str, Any], render_pages: bool,
               ocr_scanned_pages: bool = False) -> str:
    sections = []
    with fitz.open(source) as pdf:
        for number, page in enumerate(pdf, start=1):
            tables = []
            try:
                tables = [table for table in page.find_tables().tables
                          if table.row_count >= 2 and table.col_count >= 2]
            except Exception as exc:
                manifest["errors"].append({"file": str(source), "page": number,
                                           "type": type(exc).__name__, "message": f"Table detection failed: {exc}"})
            if tables:
                table_rects = [fitz.Rect(table.bbox) for table in tables]
                text_blocks = []
                for block in page.get_text("blocks", sort=True):
                    if block[6] != 0:
                        continue
                    block_rect = fitz.Rect(block[:4])
                    if not any(rect.contains(block_rect) for rect in table_rects):
                        text_blocks.append(block[4].strip())
                page_text = "\n\n".join(text_blocks)
            else:
                page_text = page.get_text(sort=True).strip()
            lines = [f"<!-- PAGE:{number} -->", page_text]
            manifest["page_map"].append({"page": number, "marker": f"<!-- PAGE:{number} -->"})
            for table in tables:
                lines.append(markdown_table(table.extract()))
            manifest.setdefault("table_count", 0)
            manifest["table_count"] += len(tables)
            for index, image in enumerate(page.get_images(full=True), start=1):
                xref = image[0]
                asset = pdf.extract_image(xref)
                asset_id = f"page-{number:04d}-image-{index:03d}"
                filename = f"{asset_id}.{asset['ext']}"
                (images_dir / filename).write_bytes(asset["image"])
                path = f"images/{filename}"
                nearby = ""
                image_rects = page.get_image_rects(xref)
                text_blocks = [block for block in page.get_text("blocks") if block[6] == 0 and block[4].strip()]
                if image_rects and text_blocks:
                    center = fitz.Rect(image_rects[0]).tl + (fitz.Rect(image_rects[0]).br - fitz.Rect(image_rects[0]).tl) / 2
                    nearest = min(text_blocks, key=lambda block: abs((block[0] + block[2]) / 2 - center.x)
                                  + abs((block[1] + block[3]) / 2 - center.y))
                    nearby = nearest[4].strip()[:200]
                manifest["image_assets"].append({"asset_id": asset_id, "source_doc_id": manifest["doc_id"],
                                                 "path": path, "page": number, "kind": "embedded",
                                                 "xref": xref, "nearby_text": nearby, "caption": None})
                lines.append(f"![PDF image {number}.{index}]({path})")
            if render_pages or (ocr_scanned_pages and not page_text.strip()):
                filename = f"page-{number:04d}.png"
                page.get_pixmap(matrix=fitz.Matrix(2, 2) if ocr_scanned_pages and not page_text.strip()
                                else fitz.Matrix(1, 1)).save(images_dir / filename)
                manifest["image_assets"].append({"asset_id": f"page-{number:04d}-render",
                                                 "source_doc_id": manifest["doc_id"],
                                                 "path": f"images/{filename}", "page": number,
                                                 "kind": "page_render", "nearby_text": "", "caption": None})
                if ocr_scanned_pages and not page_text.strip():
                    lines.append(f"![PDF page {number} OCR source](images/{filename})")
                    ocr = _ocr_text(images_dir / filename, manifest)
                    manifest["image_assets"][-1]["ocr_character_count"] = len(ocr)
                    if ocr:
                        lines.append(f"<!-- OCR:local page={number} review_status=unreviewed "
                                     f"image=images/{filename} -->\n{ocr}")
                    else:
                        manifest["errors"].append({"file": str(source), "page": number,
                                                   "type": "OCRNoText", "message": "Scanned page OCR returned no text"})
            sections.append("\n\n".join(part for part in lines if part))
    return "\n\n".join(sections) + "\n"


def _parse_xlsx(source: Path, manifest: dict[str, Any]) -> str:
    """Read cached cell values without running formulas or following external links."""
    workbook = load_workbook(source, read_only=True, data_only=True, keep_links=False)
    sections: list[str] = []
    manifest["sheets"] = []
    try:
        for sheet in workbook:
            rows: list[str] = [f"# 工作表：{sheet.title}"]
            included = 0
            max_row = sheet.max_row or 5000
            max_column = sheet.max_column or 100
            for row_number, cells in enumerate(sheet.iter_rows(max_row=min(max_row, 5000),
                                                                  max_col=min(max_column, 100)), 1):
                values = [str(cell.value).replace("|", "\\|").replace("\n", "<br>")
                          if cell.value is not None else "" for cell in cells]
                if any(values):
                    while values and not values[-1]:
                        values.pop()
                    rows.append(f"行 {row_number}：" + " | ".join(values))
                    included += 1
            manifest["sheets"].append({"name": sheet.title, "rows_with_values": included,
                                       "max_row": sheet.max_row, "max_column": sheet.max_column})
            if max_row > 5000 or max_column > 100:
                manifest["errors"].append({"file": str(source), "sheet": sheet.title,
                                           "type": "XlsxLimit", "message": "Only first 5000 rows and 100 columns parsed"})
            sections.append("\n".join(rows))
    finally:
        workbook.close()
    return "\n\n".join(sections) + "\n"


def _write_spreadsheet_snapshot(source: Path, target: Path, manifest: dict[str, Any]) -> None:
    """Persist a cell-level spreadsheet.json for table browse and bounded compute.

    A snapshot failure is recorded as a non-fatal error so Markdown parsing still
    succeeds even when the structured table model cannot be built.
    """
    try:
        from src.parser import spreadsheet
        sheets, warnings = spreadsheet.read_xlsx(source) if source.suffix.lower() == ".xlsx" \
            else spreadsheet.read_csv(source)
        spreadsheet.write_artifact(target, source, manifest["sha256"], sheets, warnings)
        manifest["spreadsheet"] = {"sheets": len(sheets),
                                   "cells": sum(len(row["cells"]) for sheet in sheets for row in sheet["rows"])}
        for warning in warnings:
            manifest["errors"].append({"file": str(source), "type": "SpreadsheetWarning", "message": warning})
    except Exception as exc:
        manifest["errors"].append({"file": str(source), "type": type(exc).__name__,
                                   "message": f"Spreadsheet snapshot failed: {exc}"})


def _parse_csv(source: Path, manifest: dict[str, Any]) -> str:
    """Read delimited text as a Markdown table, preserving leading zeros and text.

    Encoding follows a deterministic order: UTF-16 when a BOM is present, else
    UTF-8 (with optional BOM), falling back to GB18030. Field content is never
    evaluated, so values that start with ``=`` stay literal text.
    """
    raw = source.read_bytes()
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        encoding = "utf-16"
    else:
        encoding = "utf-8-sig"
    try:
        content = raw.decode(encoding)
    except UnicodeDecodeError:
        content = raw.decode("gb18030")
        encoding = "gb18030"
    content = content.replace("\r\n", "\n").replace("\r", "\n")
    manifest["encoding"] = encoding
    manifest["character_count"] = len(content)
    lines = content.splitlines()
    sample = "\n".join(lines[:5])
    try:
        delimiter = _csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except _csv.Error:
        delimiter = ","
    rows = [[cell or "" for cell in row] for row in _csv.reader(lines, delimiter=delimiter)]
    rows = [row for row in rows if any(row)]
    max_rows, max_cols = 5000, 100
    truncated = len(rows) > max_rows or any(len(row) > max_cols for row in rows)
    rows = [row[:max_cols] for row in rows[:max_rows]]
    manifest["csv"] = {"rows_with_values": len(rows), "columns": max((len(row) for row in rows), default=0),
                       "delimiter": delimiter}
    if truncated:
        manifest["errors"].append({"file": str(source), "type": "CsvLimit",
                                   "message": "Only first 5000 rows and 100 columns parsed"})
    if not rows:
        return ""
    return f"# CSV 数据：{source.stem}\n\n{markdown_table(rows)}\n"


@lru_cache(maxsize=1)
def _ocr_engine() -> RapidOCR:
    return RapidOCR()


def _ocr_text(path: Path, manifest: dict[str, Any]) -> str:
    try:
        result, _ = _ocr_engine()(str(path))
        return "\n".join(item[1] for item in result) if result else ""
    except Exception as exc:
        manifest["errors"].append({"file": str(path), "type": type(exc).__name__, "message": f"OCR failed: {exc}"})
        return ""


def _parse_docx(source: Path, images_dir: Path, manifest: dict[str, Any]) -> str:
    # DOCX paragraphs and built-in heading styles can be read directly from OOXML.
    # This keeps parsing independent of office software and page layout guesses.
    with zipfile.ZipFile(source) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
        word_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        q = lambda name: f"{{{word_ns}}}{name}"
        style_levels = {}
        if "word/styles.xml" in archive.namelist():
            styles_root = ElementTree.fromstring(archive.read("word/styles.xml"))
            for style_node in styles_root.iter(q("style")):
                style_id = style_node.get(q("styleId"))
                outline = style_node.find(f"{q('pPr')}/{q('outlineLvl')}")
                name_node = style_node.find(q("name"))
                name = name_node.get(q("val"), "") if name_node is not None else ""
                match = re.fullmatch(r"(?:heading|标题)\s*([1-6])", name, re.IGNORECASE)
                if outline is not None and outline.get(q("val"), "").isdigit():
                    level = int(outline.get(q("val"))) + 1
                elif match:
                    level = int(match.group(1))
                else:
                    continue
                if style_id and 1 <= level <= 6:
                    style_levels[style_id] = level
        rels_root = ElementTree.fromstring(archive.read("word/_rels/document.xml.rels"))
        relationships = {
            node.get("Id"): node.get("Target")
            for node in rels_root
            if node.get("Type", "").endswith("/image") and node.get("TargetMode") != "External"
        }
        image_data = {}
        for rel_id, target in relationships.items():
            member = posixpath.normpath(posixpath.join("word", target.lstrip("/")))
            if not member.startswith("word/"):
                raise ValueError(f"Unsafe DOCX image path: {target}")
            image_data[rel_id] = archive.read(member)
    relation_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    drawing_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    vml_ns = "urn:schemas-microsoft-com:vml"
    blocks = []
    extracted: dict[str, tuple[str, str]] = {}
    paragraph_index = 0
    recent_text = ""
    manifest["paragraphs"] = []
    manifest["table_count"] = 0

    def render_paragraph(paragraph: ElementTree.Element, table_index: int | None = None,
                         row_index: int | None = None, column_index: int | None = None) -> str:
        nonlocal paragraph_index, recent_text
        paragraph_index += 1
        paragraph_id = f"p-{paragraph_index:05d}"
        value = "".join(node.text or "" for node in paragraph.iter(q("t"))).strip()
        ids = [node.get(f"{{{relation_ns}}}embed") for node in paragraph.iter(f"{{{drawing_ns}}}blip")]
        ids += [node.get(f"{{{relation_ns}}}id") for node in paragraph.iter(f"{{{vml_ns}}}imagedata")]
        ids = [rel_id for rel_id in ids if rel_id in image_data]
        if not value and not ids:
            return ""
        manifest["paragraphs"].append({"paragraph_id": paragraph_id, "table_index": table_index,
                                       "row": row_index, "column": column_index})
        style_node = paragraph.find(f"{q('pPr')}/{q('pStyle')}")
        style = style_node.get(q("val"), "") if style_node is not None else ""
        level = style_levels.get(style)
        if level is None and style.startswith("Heading") and style[7:].isdigit():
            level = max(1, min(6, int(style[7:])))
        if level is not None:
            value = f"{'#' * level} {value}"
        parts = [] if table_index is not None else [f"<!-- paragraph_id:{paragraph_id} -->"]
        if value:
            parts.append(value)
        for rel_id in ids:
            if rel_id not in extracted:
                data = image_data[rel_id]
                with Image.open(BytesIO(data)) as image:
                    image.verify()
                with Image.open(BytesIO(data)) as image:
                    image_format = image.format
                    width, height = image.size
                extension = {"JPEG": "jpg", "TIFF": "tiff"}.get(image_format, image_format.lower())
                asset_id = f"docx-image-{len(extracted) + 1:03d}"
                filename = f"{asset_id}.{extension}"
                asset_path = images_dir / filename
                asset_path.write_bytes(data)
                asset = {"asset_id": asset_id, "source_doc_id": manifest["doc_id"],
                         "path": f"images/{filename}", "kind": "embedded", "relationship_id": rel_id,
                         "paragraph_id": paragraph_id, "nearby_text": (value or recent_text)[:200],
                         "caption": None, "format": image_format, "width": width, "height": height}
                ocr = _ocr_text(asset_path, manifest)
                asset["ocr_character_count"] = len(ocr)
                manifest["image_assets"].append(asset)
                extracted[rel_id] = (asset["path"], ocr)
            path, ocr = extracted[rel_id]
            parts.append(f"![DOCX image {rel_id}]({path})")
            if ocr:
                parts.append(f"<!-- OCR:local image={path} -->\n{ocr}")
        if value:
            recent_text = value
        return "\n\n".join(parts)

    body = root.find(q("body"))
    if body is None:
        return ""
    for child in body:
        if child.tag == q("p"):
            rendered = render_paragraph(child)
            if rendered:
                blocks.append(rendered)
        elif child.tag == q("tbl"):
            table_index = manifest["table_count"] + 1
            rows = []
            for row_index, row in enumerate(child.findall(q("tr")), start=1):
                cells = []
                for column_index, cell in enumerate(row.findall(q("tc")), start=1):
                    properties = cell.find(q("tcPr"))
                    span = properties.find(q("gridSpan")) if properties is not None else None
                    merge = properties.find(q("vMerge")) if properties is not None else None
                    colspan = int(span.get(q("val"), "1")) if span is not None else 1
                    vmerge = (merge.get(q("val"), "continue") if merge is not None else None)
                    contents = [render_paragraph(paragraph, table_index, row_index, column_index)
                                for paragraph in cell.findall(q("p"))]
                    cells.append({"text": "\n".join(value for value in contents if value),
                                  "colspan": colspan, "vmerge": vmerge})
                if cells:
                    rows.append(cells)
            if rows:
                manifest["table_count"] += 1
                anchors = [f"<!-- paragraph_id:{item['paragraph_id']} table:{table_index} "
                           f"row:{item['row']} column:{item['column']} -->"
                           for item in manifest["paragraphs"] if item["table_index"] == table_index]
                blocks.append("\n".join(anchors) + "\n\n" + docx_table(rows))
    manifest["paragraph_count"] = len(manifest["paragraphs"])
    return "\n\n".join(blocks) + "\n" if blocks else ""


def _parse_txt(source: Path, manifest: dict[str, Any]) -> str:
    raw = source.read_bytes()
    match = from_bytes(raw).best()
    if match is None:
        raise ValueError("Could not detect text encoding")
    content = str(match)
    manifest["encoding"] = match.encoding
    manifest["character_count"] = len(content)
    return content.replace("\r\n", "\n").replace("\r", "\n")


def _parse_image(source: Path, images_dir: Path, manifest: dict[str, Any]) -> str:
    try:
        with Image.open(source) as image:
            image.verify()
        with Image.open(source) as image:
            image_format = image.format
            width, height = image.size
    except (UnidentifiedImageError, OSError) as exc:
        raise ValueError("File content is not a valid supported image") from exc
    if image_format not in {"PNG", "JPEG", "GIF", "BMP", "WEBP", "TIFF"}:
        raise ValueError(f"Unsupported image content format: {image_format}")
    extension = {"JPEG": "jpg", "TIFF": "tiff"}.get(image_format, image_format.lower())
    filename = f"source.{extension}"
    (images_dir / filename).write_bytes(source.read_bytes())
    manifest["source_type"] = image_format.lower()
    ocr = _ocr_text(images_dir / filename, manifest)
    manifest["image_assets"].append({"asset_id": "source", "source_doc_id": manifest["doc_id"],
                                     "path": f"images/{filename}", "kind": "source",
                                     "nearby_text": "", "caption": None, "format": image_format,
                                     "width": width, "height": height, "ocr_character_count": len(ocr)})
    markdown = f"![{source.stem}](images/{filename})\n"
    if ocr:
        markdown += f"\n<!-- OCR:local image=images/{filename} -->\n{ocr}\n"
    return markdown
